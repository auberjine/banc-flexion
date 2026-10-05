# -*- coding: utf-8 -*-
"""
Nomenclature au format Excel, pour les ateliers qui ne lisent pas le Markdown.

    python nomenclature_xlsx.py

Lit NOMENCLATURE.md (produit par nomenclature.py) et ecrit
out/NOMENCLATURE.xlsx : un onglet par tableau, un onglet pour les ordres de
montage, un onglet « Lisez-moi » pour les exigences de commande. Les masses
deviennent des nombres et les masses totales des formules (qte x masse u.),
avec leur somme en bas : le classeur se recalcule si l'on change une quantite.
"""

import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SOURCE = os.path.join(HERE, "NOMENCLATURE.md")
CIBLE = os.path.join(HERE, "out", "NOMENCLATURE.xlsx")

POLICE = "Arial"
F_TITRE = Font(name=POLICE, size=14, bold=True)
F_ENTETE = Font(name=POLICE, size=10, bold=True, color="FFFFFF")
F_NORMAL = Font(name=POLICE, size=10)
F_GRAS = Font(name=POLICE, size=10, bold=True)
F_NOTE = Font(name=POLICE, size=9, italic=True, color="555555")
FOND_ENTETE = PatternFill("solid", fgColor="44546A")
FOND_TOTAL = PatternFill("solid", fgColor="E7E6E6")
FIN = Side(style="thin", color="A6A6A6")
CADRE = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)
HAUT_GAUCHE = Alignment(vertical="top", wrap_text=True)

# onglets : titre de section Markdown -> nom court (31 caracteres au plus)
NOMS = {
    "Pieces fabriquees": "Pieces fabriquees",
    "Pieces du commerce": "Pieces du commerce",
    "Visserie et petites pieces du commerce": "Visserie",
    "Pieces planes a decouper": "DXF a decouper",
    "Reglages et constantes d'essai": "Reglages",
}
LARGEUR_MAX = 70
# masses unitaires exactes du modele 3D, par designation : le Markdown les
# arrondit a 0,01 kg, ce qui fausse les totaux des pieces en plusieurs exemplaires
MASSES = {}


def sans_md(t):
    """Retire le gras et le code Markdown d'une cellule."""
    t = re.sub(r"\*\*(.*?)\*\*", r"\1", t)
    t = re.sub(r"`(.*?)`", r"\1", t)
    return t.strip()


def masse(t):
    """'0.86 kg' -> 0.86 ; sinon None."""
    m = re.fullmatch(r"\s*([0-9]+(?:[.,][0-9]+)?)\s*kg\s*", t)
    return float(m.group(1).replace(",", ".")) if m else None


def nombre(t):
    m = re.fullmatch(r"\s*([0-9]+)\s*", t)
    return int(m.group(1)) if m else None


def lire_sections(texte):
    """[(titre, [lignes])] dans l'ordre du document ; titre None pour l'en-tete."""
    sections, titre, lignes = [], None, []
    for L in texte.splitlines():
        if L.startswith("## "):
            sections.append((titre, lignes))
            titre, lignes = L[3:].strip(), []
        elif L.startswith("# "):
            continue
        else:
            lignes.append(L)
    sections.append((titre, lignes))
    return sections


def tables_et_texte(lignes):
    """Separe les tableaux Markdown du texte libre d'une section."""
    tables, texte, courante = [], [], None
    for L in lignes:
        if L.startswith("|"):
            cellules = [c.strip() for c in L.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", c) for c in cellules):
                continue
            if courante is None:
                courante = []
                tables.append(courante)
            courante.append(cellules)
        else:
            courante = None
            if L.strip():
                texte.append(L.strip())
    return tables, texte


def largeurs(ws, ncol):
    for j in range(1, ncol + 1):
        lettre = get_column_letter(j)
        lmax = max((len(str(c.value)) for c in ws[lettre] if c.value is not None), default=8)
        ws.column_dimensions[lettre].width = min(max(8, lmax + 2), LARGEUR_MAX)


def ecrire_table(ws, ligne0, table):
    """Ecrit un tableau a partir de ligne0 ; rend la ligne suivante libre."""
    entete, corps = table[0], table[1:]
    for j, h in enumerate(entete, 1):
        c = ws.cell(row=ligne0, column=j, value=sans_md(h))
        c.font, c.fill, c.border = F_ENTETE, FOND_ENTETE, CADRE
        c.alignment = Alignment(vertical="center", wrap_text=True)
    noms = [sans_md(h).lower() for h in entete]
    j_qte = noms.index("qte") + 1 if "qte" in noms else None
    j_mu = noms.index("masse u.") + 1 if "masse u." in noms else None
    j_mt = noms.index("masse tot.") + 1 if "masse tot." in noms else None
    i = ligne0
    for rangee in corps:
        i += 1
        for j, v in enumerate(rangee, 1):
            v = sans_md(v)
            c = ws.cell(row=i, column=j)
            if j == j_qte and nombre(v) is not None:
                c.value = nombre(v)
            elif j in (j_mu, j_mt) and masse(v) is not None:
                c.value = masse(v)
                if j == j_mu and sans_md(rangee[1]) in MASSES:
                    c.value = MASSES[sans_md(rangee[1])]
                c.number_format = '0.00 "kg"'
            else:
                c.value = v
            c.font, c.border, c.alignment = F_NORMAL, CADRE, HAUT_GAUCHE
        # masse totale = qte x masse unitaire : une formule, pas un nombre recopie
        if j_qte and j_mu and j_mt:
            q, mu = ws.cell(row=i, column=j_qte).value, ws.cell(row=i, column=j_mu).value
            if isinstance(q, int) and isinstance(mu, float):
                c = ws.cell(row=i, column=j_mt)
                c.value = "=%s%d*%s%d" % (get_column_letter(j_qte), i, get_column_letter(j_mu), i)
                c.number_format = '0.00 "kg"'
    # somme des masses totales en pied de tableau
    if j_mt and i > ligne0:
        i += 1
        lt = get_column_letter(j_mt)
        ws.cell(row=i, column=j_mt - 1, value="total").font = F_GRAS
        c = ws.cell(row=i, column=j_mt, value="=SUM(%s%d:%s%d)" % (lt, ligne0 + 1, lt, i - 1))
        c.font, c.number_format = F_GRAS, '0.00 "kg"'
        for j in range(1, len(entete) + 1):
            ws.cell(row=i, column=j).fill = FOND_TOTAL
            ws.cell(row=i, column=j).border = CADRE
    return i + 2


def main():
    texte = open(SOURCE, encoding="utf-8").read()
    import json
    fm = os.path.join(HERE, "out", "masses.json")
    if os.path.isfile(fm):
        d = json.load(open(fm, encoding="utf-8"))
        for x in (d if isinstance(d, list) else d.get("pieces", [])):
            MASSES[x["designation"]] = round(x["masse_kg"], 3)
    titre_doc = next((L[2:].strip() for L in texte.splitlines() if L.startswith("# ")), "Nomenclature")
    wb = Workbook()
    lisez = wb.active
    lisez.title = "Lisez-moi"
    lisez["A1"] = titre_doc
    lisez["A1"].font = F_TITRE
    lisez["A2"] = ("Genere depuis NOMENCLATURE.md par nomenclature_xlsx.py : modifier le modele "
                   "(params.py) et relancer make.py plutot que ce classeur.")
    lisez["A2"].font = F_NOTE
    ligne_lisez = 4
    montage = None
    for titre, lignes in lire_sections(texte):
        tables, libre = tables_et_texte(lignes)
        if titre in NOMS and tables:
            ws = wb.create_sheet(NOMS[titre])
            ws["A1"] = titre
            ws["A1"].font = F_TITRE
            i = 3
            for t in tables:
                i = ecrire_table(ws, i, t)
            for L in libre:
                ws.cell(row=i, column=1, value=sans_md(L)).font = F_NOTE
                i += 1
            ws.freeze_panes = "A4"
            largeurs(ws, max(len(t[0]) for t in tables))
        elif titre is not None and not tables and any(re.match(r"\d+\. ", L) for L in libre):
            # ordres de montage : etapes numerotees, un onglet pour toutes
            if montage is None:
                montage = wb.create_sheet("Montage")
                montage.column_dimensions["A"].width = 6
                montage.column_dimensions["B"].width = 120
                im = 1
            montage.cell(row=im, column=1, value=titre).font = F_TITRE
            im += 2
            for L in libre:
                m = re.match(r"(\d+)\. (.*)", L)
                if m:
                    montage.cell(row=im, column=1, value=int(m.group(1))).font = F_GRAS
                    c = montage.cell(row=im, column=2, value=sans_md(m.group(2)))
                else:
                    c = montage.cell(row=im, column=2, value=sans_md(L))
                    c.font = F_NOTE
                c.alignment = HAUT_GAUCHE
                if c.font != F_NOTE:
                    c.font = F_NORMAL
                im += 1
            im += 1
        else:
            # en-tete du document et exigences de commande : sur l'onglet Lisez-moi
            if titre:
                lisez.cell(row=ligne_lisez, column=1, value=titre).font = F_GRAS
                ligne_lisez += 1
            for L in libre:
                c = lisez.cell(row=ligne_lisez, column=1, value=sans_md(L.lstrip("-* ")))
                c.font, c.alignment = F_NORMAL, HAUT_GAUCHE
                ligne_lisez += 1
            for t in tables:
                ligne_lisez = ecrire_table(lisez, ligne_lisez, t)
            ligne_lisez += 1
    lisez.column_dimensions["A"].width = 120
    os.makedirs(os.path.dirname(CIBLE), exist_ok=True)
    # openpyxl n'ecrit pas les valeurs des formules : Excel recalcule a l'ouverture
    wb.calculation.fullCalcOnLoad = True
    wb.save(CIBLE)
    print(CIBLE, ", ".join(wb.sheetnames))


if __name__ == "__main__":
    main()
